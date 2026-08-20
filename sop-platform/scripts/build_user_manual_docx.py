from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCES = [
    ROOT / "docs" / "18_SOP平台完整使用说明_20260820.md",
    ROOT / "docs" / "19_Windows_APP安装运行说明_20260820.md",
]


def apply_run_font(run, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "KaiTi")
    run.font.size = Pt(14)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, text: str, bold: bool = False) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        marked = (part.startswith("`") and part.endswith("`")) or (part.startswith("**") and part.endswith("**"))
        content = part[1:-1] if part.startswith("`") else part[2:-2] if part.startswith("**") else part
        apply_run_font(paragraph.add_run(content), bold or marked)


def build(source: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Pt(56)
    section.left_margin = section.right_margin = Pt(62)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "KaiTi")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)

    in_code = False
    for raw in source.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if not line:
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(18)
            add_inline(paragraph, line[2:], True)
        elif line.startswith("## "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, line[3:], True)
        elif re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            if in_code:
                paragraph.paragraph_format.left_indent = Pt(18)
            add_inline(paragraph, line)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(6)
        for run in paragraph.runs:
            apply_run_font(run, run.bold)

    output = source.with_suffix(".docx")
    document.save(output)
    return output


def main() -> None:
    for source in SOURCES:
        print(build(source))


if __name__ == "__main__":
    main()
